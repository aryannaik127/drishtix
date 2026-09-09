import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(*(int(hex_str[i:i+2], 16) for i in (0, 2, 4)))

# Color Palette
C_BG_WHITE = hex_to_rgb('FFFFFF')
C_NAVY_TITLE = hex_to_rgb('1E3A8A') # SIH Navy (Blue 900)
C_DARK_TEXT = hex_to_rgb('0F172A')  # Slate 900
C_BODY_TEXT = hex_to_rgb('334155')  # Slate 700
C_MUTED_TEXT = hex_to_rgb('64748B') # Slate 500
C_FOOTER_BLUE = hex_to_rgb('1E40AF') # Blue 800

# Theme Colors
C_PURPLE = hex_to_rgb('6B46C1')
C_PURPLE_DARK = hex_to_rgb('553C9A')
C_PURPLE_BG = hex_to_rgb('F5F3FF')
C_PURPLE_BORDER = hex_to_rgb('DDD6FE')

C_BLUE = hex_to_rgb('1D4ED8')
C_BLUE_LIGHT = hex_to_rgb('2563EB')
C_BLUE_BG = hex_to_rgb('EFF6FF')
C_BLUE_BORDER = hex_to_rgb('BFDBFE')

C_TEAL = hex_to_rgb('0D9488')
C_TEAL_DARK = hex_to_rgb('0F766E')
C_TEAL_BG = hex_to_rgb('F0FDFA')
C_TEAL_BORDER = hex_to_rgb('99F6E4')

C_ORANGE = hex_to_rgb('D97706')
C_ORANGE_DARK = hex_to_rgb('EA580C')
C_ORANGE_BG = hex_to_rgb('FFFBEB')
C_ORANGE_BORDER = hex_to_rgb('FDE68A')

C_RED = hex_to_rgb('DC2626')
C_RED_DARK = hex_to_rgb('B91C1C')
C_RED_BG = hex_to_rgb('FEF2F2')
C_RED_BORDER = hex_to_rgb('FECACA')

C_GREEN = hex_to_rgb('16A34A')
C_GREEN_DARK = hex_to_rgb('15803D')
C_GREEN_BG = hex_to_rgb('F0FDF4')
C_GREEN_BORDER = hex_to_rgb('BBF7D0')

FONT_TITLE = 'Arial'
FONT_BODY = 'Segoe UI'

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    assets_dir = r'c:\Users\FALCON JNB\.gemini\antigravity-ide\scratch\drishtix\extracted_assets'
    sih_logo_path = os.path.join(assets_dir, 'sih_logo.png')
    drishtix_logo_path = os.path.join(assets_dir, 'drishtix_logo.png')
    brain_graphic_path = os.path.join(assets_dir, 'slide1_brain_bulb.png')
    proto_screen_path = os.path.join(assets_dir, 'prototype_screen.png')

    def add_header(slide, title_text, subtitle_text=None, page_num=None, show_drishtix_logo=True):
        # Top Logos
        if show_drishtix_logo and os.path.exists(drishtix_logo_path):
            slide.shapes.add_picture(drishtix_logo_path, Inches(0.4), Inches(0.16), width=Inches(1.5))
        
        if os.path.exists(sih_logo_path):
            slide.shapes.add_picture(sih_logo_path, Inches(10.8), Inches(0.12), width=Inches(2.15))

        # Title Box
        if title_text:
            title_box = slide.shapes.add_textbox(
                Inches(1.8 if show_drishtix_logo else 0.8), 
                Inches(0.16 if not subtitle_text else 0.14), 
                Inches(8.8), 
                Inches(0.65 if not subtitle_text else 0.45)
            )
            tf = title_box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = title_text
            p.font.name = FONT_TITLE
            p.font.size = Pt(21 if subtitle_text else 26)
            p.font.bold = True
            p.font.color.rgb = C_NAVY_TITLE
            if not show_drishtix_logo or not subtitle_text:
                p.alignment = PP_ALIGN.CENTER

            if subtitle_text:
                p2 = tf.add_paragraph()
                p2.text = subtitle_text
                p2.font.name = FONT_BODY
                p2.font.size = Pt(13.5)
                p2.font.italic = True
                p2.font.bold = True
                p2.font.color.rgb = hex_to_rgb('1E40AF')

        # Footer
        if page_num is not None:
            footer_box = slide.shapes.add_textbox(Inches(4.0), Inches(7.12), Inches(5.33), Inches(0.3))
            ftf = footer_box.text_frame
            ftf.word_wrap = False
            ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
            p = ftf.paragraphs[0]
            p.text = "@SIH Idea submission- Template"
            p.font.name = FONT_BODY
            p.font.size = Pt(10.5)
            p.font.color.rgb = C_FOOTER_BLUE
            p.alignment = PP_ALIGN.CENTER

            # Page Number
            pnum_box = slide.shapes.add_textbox(Inches(12.5), Inches(7.12), Inches(0.5), Inches(0.3))
            ptf = pnum_box.text_frame
            ptf.word_wrap = False
            ptf.margin_left = ptf.margin_top = ptf.margin_right = ptf.margin_bottom = 0
            pp = ptf.paragraphs[0]
            pp.text = str(page_num)
            pp.font.name = FONT_BODY
            pp.font.size = Pt(11.5)
            pp.font.bold = True
            pp.font.color.rgb = C_FOOTER_BLUE
            pp.alignment = PP_ALIGN.RIGHT

    # ==========================================
    # SLIDE 1: Title & Team Registration Slide
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    
    # SIH 2026 Logo Top Right
    if os.path.exists(sih_logo_path):
        slide1.shapes.add_picture(sih_logo_path, Inches(10.6), Inches(0.18), width=Inches(2.35))

    # Main Title Header
    s1_title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.5), Inches(0.8))
    s1_tf = s1_title_box.text_frame
    s1_tf.word_wrap = True
    s1_tf.margin_left = s1_tf.margin_top = s1_tf.margin_right = s1_tf.margin_bottom = 0
    p = s1_tf.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2026"
    p.font.name = FONT_TITLE
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_TITLE
    p.alignment = PP_ALIGN.CENTER

    # Brain / Bulb Graphic on Right
    if os.path.exists(brain_graphic_path):
        slide1.shapes.add_picture(brain_graphic_path, Inches(6.6), Inches(1.3), width=Inches(5.8))

    # Left Fields Box (Editable text)
    fields_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(6.0), Inches(4.8))
    ftf = fields_box.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0

    items_s1 = [
        "Problem Statement ID \u2013",
        "Problem Statement Title -",
        "Theme -",
        "PS Category - Software/Hardware",
        "Team ID -",
        "Team Name (Registered on portal)"
    ]

    for idx, item in enumerate(items_s1):
        p = ftf.paragraphs[0] if idx == 0 else ftf.add_paragraph()
        p.text = f"\u2022  {item}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb('0F172A')
        p.space_after = Pt(22)

    # ==========================================
    # SLIDE 2: Idea, Problem, Solution, Benefits & Demo
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "DRISHTIX \u2013 INTELLIGENT BORDER SURVEILLANCE PLATFORM", 
               "Turning Existing CCTV infrastructure into intelligent Actionable intelligence", page_num=2)

    # Row 1: 5 Cards
    card_w = Inches(2.38)
    card_gap = Inches(0.12)
    start_x = Inches(0.4)
    top_y = Inches(1.32)
    card_h = Inches(2.68)

    cards_r1 = [
        {
            "num": "1. OUR IDEA", "color": C_PURPLE, "bg": C_PURPLE_BG, "border": C_PURPLE_BORDER, "icon": "\U0001F4A1",
            "type": "text",
            "content": "Create a software-defined AI platform that converts existing CCTV cameras into an intelligent surveillance network without requiring expensive hardware."
        },
        {
            "num": "2. THE PROBLEM", "color": C_BLUE_LIGHT, "bg": C_BLUE_BG, "border": C_BLUE_BORDER, "icon": "\U0001F9E9",
            "type": "bullets",
            "bullets": [
                "Existing CCTV is passive and requires constant manual monitoring.",
                "High false alarms with conventional motion detection.",
                "Advanced features like FRS, ANPR, intrusion detection require costly hardware.",
                "Difficult & expensive to deploy at large scale, especially in remote border areas."
            ]
        },
        {
            "num": "3. OUR SOLUTION (DRISHTIX)", "color": C_TEAL, "bg": C_TEAL_BG, "border": C_TEAL_BORDER, "icon": "\U0001F3AF",
            "type": "grid_solution",
            "title": "AI-Powered Video Analytics Platform",
            "items": [
                "Human Detection & Tracking",
                "Vehicle Detection & Classification",
                "Face Detection",
                "ANPR (Plate Recognition)",
                "Virtual Fence Intrusion Detection",
                "Suspicious Activity Detection",
                "Night-Time Movement Detection",
                "Real-Time Alerts & Event Logging"
            ]
        },
        {
            "num": "4. KEY BENEFITS", "color": C_ORANGE_DARK, "bg": C_ORANGE_BG, "border": C_ORANGE_BORDER, "icon": "\U0001F6E1",
            "type": "bullets",
            "bullets": [
                "Works with existing CCTV infrastructure",
                "Eliminates dependency on expensive hardware",
                "Real-time alerts for quick response",
                "Improved situational awareness for border security",
                "Scalable, cost-effective for remote locations",
                "Seamless integration with Command & Control systems"
            ]
        },
        {
            "num": "5. TECHNOLOGIES USED", "color": C_RED, "bg": C_RED_BG, "border": C_RED_BORDER, "icon": "\u2699",
            "type": "tech_bullets",
            "bullets": [
                "Python, FastAPI",
                "OpenCV",
                "YOLOv8",
                "Deep Learning (Autoencoders, LSTM)",
                "OCR / ANPR",
                "PostgreSQL / SQLite",
                "React Dashboard"
            ]
        }
    ]

    for idx, c in enumerate(cards_r1):
        cx = start_x + idx * (card_w + card_gap)
        
        # Outer Card Container Shape
        card_shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, top_y, card_w, card_h)
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = c['bg']
        card_shape.line.color.rgb = c['border']
        card_shape.line.width = Pt(1.2)

        # Top Header Banner
        banner = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.04), top_y + Inches(0.04), card_w - Inches(0.08), Inches(0.38))
        banner.fill.solid()
        banner.fill.fore_color.rgb = c['color']
        banner.line.color.rgb = c['color']
        btf = banner.text_frame
        btf.margin_left = btf.margin_right = btf.margin_top = btf.margin_bottom = 0
        bp = btf.paragraphs[0]
        bp.text = f"{c.get('icon', '')} {c['num']}"
        bp.font.name = FONT_TITLE
        bp.font.size = Pt(9.5)
        bp.font.bold = True
        bp.font.color.rgb = C_BG_WHITE
        bp.alignment = PP_ALIGN.CENTER

        # Content Inside
        content_box = slide2.shapes.add_textbox(cx + Inches(0.08), top_y + Inches(0.44), card_w - Inches(0.16), card_h - Inches(0.48))
        ctf = content_box.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0

        if c['type'] == 'text':
            p = ctf.paragraphs[0]
            p.text = c['content']
            p.font.name = FONT_BODY
            p.font.size = Pt(10.5)
            p.font.color.rgb = C_DARK_TEXT
            p.space_before = Pt(8)
            p.alignment = PP_ALIGN.LEFT
        elif c['type'] == 'bullets':
            for bidx, btext in enumerate(c['bullets']):
                p = ctf.paragraphs[0] if bidx == 0 else ctf.add_paragraph()
                p.text = f"\u2022 {btext}"
                p.font.name = FONT_BODY
                p.font.size = Pt(8.5)
                p.font.color.rgb = C_DARK_TEXT
                p.space_after = Pt(2.5)
        elif c['type'] == 'grid_solution':
            p = ctf.paragraphs[0]
            p.text = c['title']
            p.font.name = FONT_TITLE
            p.font.size = Pt(8.8)
            p.font.bold = True
            p.font.color.rgb = C_TEAL_DARK
            p.alignment = PP_ALIGN.CENTER
            p.space_after = Pt(2)

            for sitem in c['items']:
                p = ctf.add_paragraph()
                p.text = f"\u2022 {sitem}"
                p.font.name = FONT_BODY
                p.font.size = Pt(7.8)
                p.font.color.rgb = C_DARK_TEXT
                p.space_after = Pt(1)
        elif c['type'] == 'tech_bullets':
            for bidx, btext in enumerate(c['bullets']):
                p = ctf.paragraphs[0] if bidx == 0 else ctf.add_paragraph()
                p.text = f"\u2022 {btext}"
                p.font.name = FONT_BODY
                p.font.size = Pt(8.5)
                p.font.color.rgb = C_DARK_TEXT
                p.space_after = Pt(2.5)

    # Row 2: Workflow (Left) and Demo View (Right)
    bot_y = Inches(4.1)
    bot_h = Inches(2.95)
    
    # Left Card: Solution Prototype - Workflow (Width 7.1 in)
    wf_w = Inches(7.1)
    wf_shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), bot_y, wf_w, bot_h)
    wf_shape.fill.solid()
    wf_shape.fill.fore_color.rgb = hex_to_rgb('F8FAFC')
    wf_shape.line.color.rgb = C_BLUE_BORDER
    wf_shape.line.width = Pt(1.2)

    # Banner for Workflow
    wfb = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.44), bot_y + Inches(0.04), wf_w - Inches(0.08), Inches(0.36))
    wfb.fill.solid()
    wfb.fill.fore_color.rgb = C_BLUE_LIGHT
    wfb.line.color.rgb = C_BLUE_LIGHT
    wb_tf = wfb.text_frame
    wb_p = wb_tf.paragraphs[0]
    wb_p.text = "6. SOLUTION PROTOTYPE \u2013 WORKFLOW"
    wb_p.font.name = FONT_TITLE
    wb_p.font.size = Pt(10)
    wb_p.font.bold = True
    wb_p.font.color.rgb = C_BG_WHITE
    wb_p.alignment = PP_ALIGN.CENTER

    # 5 Workflow Stage Boxes
    stages = [
        {"title": "INPUT", "sub": "Existing CCTV\n(IP / RTSP Feeds)", "color": hex_to_rgb('6366F1')},
        {"title": "PREPROCESSING", "sub": "Frame Extraction\nResize \u2022 Noise Reduc.\nNormalization", "color": hex_to_rgb('0284C7')},
        {"title": "AI ANALYTICS", "sub": "Detection & Track\nFeature Extraction\nAnomaly Detection\nRule & Event Anal.", "color": hex_to_rgb('10B981')},
        {"title": "OUTPUT", "sub": "Real-time Alerts\nEvent Logging\nSnapshots / Clips\nRisk Score", "color": hex_to_rgb('F59E0B')},
        {"title": "DASHBOARD", "sub": "Live Monitoring\nEvent Tagging\nSearch & Filter\nReports & Analytics", "color": hex_to_rgb('EC4899')}
    ]

    st_start_x = Inches(0.55)
    st_w = Inches(1.18)
    st_gap = Inches(0.12)
    st_y = bot_y + Inches(0.52)
    st_h = Inches(2.28)

    for sidx, st in enumerate(stages):
        sx = st_start_x + sidx * (st_w + st_gap)
        sbox = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, st_y, st_w, st_h)
        sbox.fill.solid()
        sbox.fill.fore_color.rgb = C_BG_WHITE
        sbox.line.color.rgb = st['color']
        sbox.line.width = Pt(1.5)

        # Stage Header
        shb = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx + Inches(0.04), st_y + Inches(0.04), st_w - Inches(0.08), Inches(0.38))
        shb.fill.solid()
        shb.fill.fore_color.rgb = st['color']
        shb.line.color.rgb = st['color']
        sbtf = shb.text_frame
        sbtf.margin_left = sbtf.margin_right = sbtf.margin_top = sbtf.margin_bottom = 0
        p = sbtf.paragraphs[0]
        p.text = st['title']
        p.font.name = FONT_TITLE
        p.font.size = Pt(7.8)
        p.font.bold = True
        p.font.color.rgb = C_BG_WHITE
        p.alignment = PP_ALIGN.CENTER

        # Stage Body
        sbtext = slide2.shapes.add_textbox(sx + Inches(0.03), st_y + Inches(0.48), st_w - Inches(0.06), st_h - Inches(0.52))
        stf = sbtext.text_frame
        stf.word_wrap = True
        stf.margin_left = stf.margin_right = stf.margin_top = stf.margin_bottom = 0
        p = stf.paragraphs[0]
        p.text = st['sub']
        p.font.name = FONT_BODY
        p.font.size = Pt(7.5)
        p.font.color.rgb = C_DARK_TEXT
        p.alignment = PP_ALIGN.CENTER

        # Arrow between stages
        if sidx < len(stages) - 1:
            arr_box = slide2.shapes.add_textbox(sx + st_w - Inches(0.02), st_y + Inches(0.9), Inches(0.16), Inches(0.3))
            atf = arr_box.text_frame
            atf.margin_left = atf.margin_right = atf.margin_top = atf.margin_bottom = 0
            p = atf.paragraphs[0]
            p.text = "\u25B6"
            p.font.name = FONT_BODY
            p.font.size = Pt(9)
            p.font.bold = True
            p.font.color.rgb = hex_to_rgb('94A3B8')

    # Right Card: Prototype Screen (Demo View) (Width 5.3 in)
    demo_x = Inches(7.62)
    demo_w = Inches(5.31)
    demo_shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, demo_x, bot_y, demo_w, bot_h)
    demo_shape.fill.solid()
    demo_shape.fill.fore_color.rgb = hex_to_rgb('1E1B4B')
    demo_shape.line.color.rgb = C_PURPLE
    demo_shape.line.width = Pt(1.2)

    # Demo Banner
    db = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, demo_x + Inches(0.04), bot_y + Inches(0.04), demo_w - Inches(0.08), Inches(0.36))
    db.fill.solid()
    db.fill.fore_color.rgb = C_PURPLE
    db.line.color.rgb = C_PURPLE
    dbtf = db.text_frame
    dbp = dbtf.paragraphs[0]
    dbp.text = "\U0001F4FA 7. PROTOTYPE SCREEN (DEMO VIEW)"
    dbp.font.name = FONT_TITLE
    dbp.font.size = Pt(10)
    dbp.font.bold = True
    dbp.font.color.rgb = C_BG_WHITE
    dbp.alignment = PP_ALIGN.CENTER

    # Demo Image inside
    if os.path.exists(proto_screen_path):
        slide2.shapes.add_picture(proto_screen_path, demo_x + Inches(0.08), bot_y + Inches(0.44), width=demo_w - Inches(0.16), height=bot_h - Inches(0.52))

    # ==========================================
    # SLIDE 3: Technical Approach & Methodology
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "TECHNICAL APPROACH", page_num=3)

    # Accent line under title
    line = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(0.82), Inches(4.33), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = C_NAVY_TITLE
    line.line.color.rgb = C_NAVY_TITLE

    # Row 1: 5 Tech Architecture Pillars
    p_w = Inches(2.38)
    p_gap = Inches(0.12)
    p_start_x = Inches(0.4)
    p_top_y = Inches(1.0)
    p_h = Inches(2.25)

    pillars = [
        {
            "num": "</> 1. PROGRAMMING", "sub": "Python", "color": C_BLUE, "bg": C_BLUE_BG, "border": C_BLUE_BORDER,
            "bullets": ["Core development", "AI/ML pipeline", "Automation", "Alert handling", "Integration & APIs"]
        },
        {
            "num": "2. COMPUTER VISION", "sub": "OpenCV", "color": C_GREEN, "bg": C_GREEN_BG, "border": C_GREEN_BORDER,
            "bullets": ["Video processing", "Frame extraction", "Image preprocessing", "Object / Face detection", "Tracking & contouring"]
        },
        {
            "num": "3. AI / MACHINE LEARNING", "sub": "Deep Learning", "color": C_PURPLE, "bg": C_PURPLE_BG, "border": C_PURPLE_BORDER,
            "bullets": ["YOLO-based detection", "Autoencoders for anomaly", "LSTM for temporal seq", "Behavior analytics"]
        },
        {
            "num": "4. ANPR & ANALYTICS", "sub": "OCR / ANPR", "color": C_ORANGE_DARK, "bg": C_ORANGE_BG, "border": C_ORANGE_BORDER,
            "bullets": ["Number plate detection", "Character recognition", "Vehicle identification", "Event classification", "Confidence scoring"]
        },
        {
            "num": "5. PLATFORM & STORAGE", "sub": "FastAPI + Database", "color": C_TEAL, "bg": C_TEAL_BG, "border": C_TEAL_BORDER,
            "bullets": ["FastAPI for backend APIs", "Event logging", "PostgreSQL / SQLite", "Data storage & retrieval", "Web dashboard"]
        }
    ]

    for idx, p in enumerate(pillars):
        px = p_start_x + idx * (p_w + p_gap)
        
        # Pillar shape
        pshape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, p_top_y, p_w, p_h)
        pshape.fill.solid()
        pshape.fill.fore_color.rgb = p['bg']
        pshape.line.color.rgb = p['border']
        pshape.line.width = Pt(1.2)

        # Header Pill
        pbanner = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px + Inches(0.04), p_top_y + Inches(0.04), p_w - Inches(0.08), Inches(0.34))
        pbanner.fill.solid()
        pbanner.fill.fore_color.rgb = p['color']
        pbanner.line.color.rgb = p['color']
        pbtf = pbanner.text_frame
        pbtf.margin_left = pbtf.margin_right = pbtf.margin_top = pbtf.margin_bottom = 0
        pbp = pbtf.paragraphs[0]
        pbp.text = p['num']
        pbp.font.name = FONT_TITLE
        pbp.font.size = Pt(8.5)
        pbp.font.bold = True
        pbp.font.color.rgb = C_BG_WHITE
        pbp.alignment = PP_ALIGN.CENTER

        # Subtitle & Bullets
        pcontent = slide3.shapes.add_textbox(px + Inches(0.06), p_top_y + Inches(0.4), p_w - Inches(0.12), p_h - Inches(0.44))
        pctf = pcontent.text_frame
        pctf.word_wrap = True
        pctf.margin_left = pctf.margin_right = pctf.margin_top = pctf.margin_bottom = 0
        
        # Subtitle
        psub = pctf.paragraphs[0]
        psub.text = p['sub']
        psub.font.name = FONT_TITLE
        psub.font.size = Pt(10)
        psub.font.bold = True
        psub.font.color.rgb = p['color']
        psub.alignment = PP_ALIGN.CENTER
        psub.space_after = Pt(2)

        # Bullets
        for bidx, btext in enumerate(p['bullets']):
            p_bullet = pctf.add_paragraph()
            p_bullet.text = f"\u2022 {btext}"
            p_bullet.font.name = FONT_BODY
            p_bullet.font.size = Pt(8.2)
            p_bullet.font.color.rgb = C_DARK_TEXT
            p_bullet.space_after = Pt(1.5)

    # Middle Banner: 5 Step Pill Chips Flow
    flow_y = Inches(3.35)
    flow_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), flow_y, Inches(12.533), Inches(0.46))
    flow_box.fill.solid()
    flow_box.fill.fore_color.rgb = hex_to_rgb('F1F5F9')
    flow_box.line.color.rgb = hex_to_rgb('CBD5E1')
    flow_box.line.width = Pt(1)

    flow_chips = [
        ("INPUT", "Existing CCTV / RTSP Feeds", C_BLUE),
        ("AI PROCESSING", "Computer Vision & Deep Learning", C_GREEN),
        ("INTELLIGENCE", "Detection, Tracking & Analysis", C_PURPLE),
        ("ALERT", "Real-Time Alerts & Notifications", C_ORANGE_DARK),
        ("RESPONSE", "Action by Security Forces", C_TEAL)
    ]

    chip_w = Inches(2.25)
    chip_gap = Inches(0.26)
    chip_start_x = Inches(0.55)

    for cidx, (c_tag, c_desc, c_col) in enumerate(flow_chips):
        cx = chip_start_x + cidx * (chip_w + chip_gap)
        c_shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, flow_y + Inches(0.04), chip_w, Inches(0.38))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = C_BG_WHITE
        c_shape.line.color.rgb = c_col
        c_shape.line.width = Pt(1)

        ctf = c_shape.text_frame
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
        cp = ctf.paragraphs[0]
        cp.text = f"{c_tag}: {c_desc}"
        cp.font.name = FONT_TITLE
        cp.font.size = Pt(6.8)
        cp.font.bold = True
        cp.font.color.rgb = c_col
        cp.alignment = PP_ALIGN.CENTER

        if cidx < len(flow_chips) - 1:
            arr_box = slide3.shapes.add_textbox(cx + chip_w, flow_y + Inches(0.04), chip_gap, Inches(0.38))
            atf = arr_box.text_frame
            atf.margin_left = atf.margin_right = atf.margin_top = atf.margin_bottom = 0
            p = atf.paragraphs[0]
            p.text = "\u2192"
            p.font.name = FONT_TITLE
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = hex_to_rgb('64748B')
            p.alignment = PP_ALIGN.CENTER

    # Row 2 (Bottom Section): 2. Methodology & Implementation (6 steps) + Demo flow
    meth_y = Inches(3.92)
    meth_h = Inches(3.1)
    
    # Left Box: Methodology (Width 8.1 in)
    meth_w = Inches(8.1)
    meth_shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), meth_y, meth_w, meth_h)
    meth_shape.fill.solid()
    meth_shape.fill.fore_color.rgb = C_BG_WHITE
    meth_shape.line.color.rgb = hex_to_rgb('CBD5E1')
    meth_shape.line.width = Pt(1.2)

    # Methodology Section Title
    m_title_box = slide3.shapes.add_textbox(Inches(0.55), meth_y + Inches(0.06), meth_w - Inches(0.3), Inches(0.35))
    mtf = m_title_box.text_frame
    mtf.margin_left = mtf.margin_top = mtf.margin_right = mtf.margin_bottom = 0
    mp = mtf.paragraphs[0]
    mp.text = "2. METHODOLOGY & IMPLEMENTATION  \u2014  DRISHTIX : From Video to Actionable Intelligence"
    mp.font.name = FONT_TITLE
    mp.font.size = Pt(10)
    mp.font.bold = True
    mp.font.color.rgb = C_NAVY_TITLE

    # 6 Steps Pipeline
    pipe_steps = [
        {"num": "1", "title": "VIDEO INPUT", "color": C_BLUE, "items": ["Capture video from existing CCTV cameras", "(IP / RTSP Feeds)"]},
        {"num": "2", "title": "PREPROCESSING", "color": C_GREEN, "items": ["Frame extraction, resize,", "normalization, noise reduction,", "enhancement"]},
        {"num": "3", "title": "AI ANALYTICS", "color": C_PURPLE, "items": ["Human / Vehicle Detection", "Face Detection, ANPR", "Tracking (DeepSORT)", "Feature Extraction"]},
        {"num": "4", "title": "EVENT ANALYSIS", "color": C_ORANGE_DARK, "items": ["Virtual Fence Intrusion", "Suspicious Activity", "Night-Time Movement", "Rule Engine & Risk Scoring"]},
        {"num": "5", "title": "ALERT ENGINE", "color": C_RED, "items": ["Real-Time Alerts", "Risk Class (Low/Med/High)", "Notifications: Email, SMS,", "WhatsApp & Dashboard"]},
        {"num": "6", "title": "DASHBOARD & LOGS", "color": C_TEAL, "items": ["Live View (Multi-Camera)", "Event Log & Snapshots", "Search, Filter & Reports", "Export & Integration APIs"]}
    ]

    step_col_w = Inches(1.22)
    step_gap = Inches(0.08)
    step_start_x = Inches(0.52)
    step_y = meth_y + Inches(0.46)
    step_h = Inches(2.5)

    for sidx, ps in enumerate(pipe_steps):
        sx = step_start_x + sidx * (step_col_w + step_gap)
        s_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, step_y, step_col_w, step_h)
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = hex_to_rgb('F8FAFC')
        s_box.line.color.rgb = ps['color']
        s_box.line.width = Pt(1.2)

        # Number circle / pill
        npill = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx + Inches(0.04), step_y + Inches(0.04), step_col_w - Inches(0.08), Inches(0.36))
        npill.fill.solid()
        npill.fill.fore_color.rgb = ps['color']
        npill.line.color.rgb = ps['color']
        ntf = npill.text_frame
        ntf.margin_left = ntf.margin_right = ntf.margin_top = ntf.margin_bottom = 0
        p = ntf.paragraphs[0]
        p.text = f"{ps['num']}. {ps['title']}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(7.2)
        p.font.bold = True
        p.font.color.rgb = C_BG_WHITE
        p.alignment = PP_ALIGN.CENTER

        # Step Text
        stext = slide3.shapes.add_textbox(sx + Inches(0.03), step_y + Inches(0.42), step_col_w - Inches(0.06), step_h - Inches(0.46))
        stf = stext.text_frame
        stf.word_wrap = True
        stf.margin_left = stf.margin_right = stf.margin_top = stf.margin_bottom = 0
        for itm_idx, itm in enumerate(ps['items']):
            p = stf.paragraphs[0] if itm_idx == 0 else stf.add_paragraph()
            p.text = f"\u2022 {itm}"
            p.font.name = FONT_BODY
            p.font.size = Pt(7.2)
            p.font.color.rgb = C_DARK_TEXT
            p.space_after = Pt(1.5)

    # Right Box: Working Prototype (Demo Flow) (Width 4.3 in)
    pf_x = Inches(8.62)
    pf_w = Inches(4.31)
    pf_shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pf_x, meth_y, pf_w, meth_h)
    pf_shape.fill.solid()
    pf_shape.fill.fore_color.rgb = hex_to_rgb('F8FAFC')
    pf_shape.line.color.rgb = C_NAVY_TITLE
    pf_shape.line.width = Pt(1.2)

    # Header
    pf_h_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pf_x + Inches(0.04), meth_y + Inches(0.04), pf_w - Inches(0.08), Inches(0.32))
    pf_h_box.fill.solid()
    pf_h_box.fill.fore_color.rgb = C_NAVY_TITLE
    pf_h_box.line.color.rgb = C_NAVY_TITLE
    pftf = pf_h_box.text_frame
    pfp = pftf.paragraphs[0]
    pfp.text = "WORKING PROTOTYPE (DEMO FLOW)"
    pfp.font.name = FONT_TITLE
    pfp.font.size = Pt(9)
    pfp.font.bold = True
    pfp.font.color.rgb = C_BG_WHITE
    pfp.alignment = PP_ALIGN.CENTER

    # Left: Vertical flow diagram
    vflow_box = slide3.shapes.add_textbox(pf_x + Inches(0.08), meth_y + Inches(0.4), Inches(2.1), meth_h - Inches(0.48))
    vtf = vflow_box.text_frame
    vtf.word_wrap = True
    vtf.margin_left = vtf.margin_top = vtf.margin_right = vtf.margin_bottom = 0
    vsteps = [
        "\U0001F4F9 Camera Feed (Live Video)",
        "\u2193",
        "\U0001F9E0 AI Detection (Objects/Faces)",
        "\u2193",
        "\u26A0\uFE0F Event Identified (Anomaly)",
        "\u2193",
        "\U0001F514 Alert Generated (Real-Time)",
        "\u2193",
        "\U0001F4BE Evidence Stored (Clip + Log)"
    ]
    for vidx, vs in enumerate(vsteps):
        p = vtf.paragraphs[0] if vidx == 0 else vtf.add_paragraph()
        p.text = vs
        p.font.name = FONT_BODY
        p.font.size = Pt(7.5) if "\u2193" not in vs else Pt(8)
        p.font.bold = True if "\u2193" not in vs else False
        p.font.color.rgb = C_NAVY_TITLE if "\u2193" not in vs else hex_to_rgb('64748B')
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(1)

    # Right: Example Alert Card
    alert_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pf_x + Inches(2.22), meth_y + Inches(0.42), Inches(1.98), meth_h - Inches(0.52))
    alert_card.fill.solid()
    alert_card.fill.fore_color.rgb = C_RED_BG
    alert_card.line.color.rgb = C_RED
    alert_card.line.width = Pt(1.5)

    ac_tf = alert_card.text_frame
    ac_tf.word_wrap = True
    ac_tf.margin_left = ac_tf.margin_right = Inches(0.06)
    ac_tf.margin_top = Inches(0.06)
    
    p = ac_tf.paragraphs[0]
    p.text = "\u26A0\uFE0F EXAMPLE ALERT"
    p.font.name = FONT_TITLE
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = C_RED
    p.alignment = PP_ALIGN.CENTER

    p2 = ac_tf.add_paragraph()
    p2.text = "VIRTUAL FENCE INTRUSION"
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(8)
    p2.font.bold = True
    p2.font.color.rgb = hex_to_rgb('991B1B')
    p2.alignment = PP_ALIGN.CENTER
    p2.space_after = Pt(4)

    alert_details = [
        "Camera: CAM-04",
        "Time: 22:14:08",
        "Date: 20-05-2026",
        "Risk Level: HIGH \U0001F534",
        "\u2714 Evidence Saved (Clip+Snap)"
    ]
    for ad in alert_details:
        p = ac_tf.add_paragraph()
        p.text = ad
        p.font.name = FONT_BODY
        p.font.size = Pt(7.5)
        p.font.bold = True if "HIGH" in ad or "Saved" in ad else False
        p.font.color.rgb = C_DARK_TEXT
        p.space_after = Pt(2)

    # ==========================================
    # SLIDE 4: Feasibility and Viability
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "FEASIBILITY AND VIABILITY", page_num=4)

    line4 = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.3), Inches(0.82), Inches(4.73), Inches(0.04))
    line4.fill.solid()
    line4.fill.fore_color.rgb = C_NAVY_TITLE
    line4.line.color.rgb = C_NAVY_TITLE

    # 3 Main Columns
    col_w = Inches(4.05)
    col_gap = Inches(0.19)
    col_start_x = Inches(0.4)
    col_top_y = Inches(1.05)
    col_h = Inches(4.9)

    cols_s4 = [
        {
            "title": "FEASIBILITY OF THE IDEA", "color": C_GREEN_DARK, "bg": C_GREEN_BG, "border": C_GREEN_BORDER, "icon": "\u2705",
            "items": [
                "Uses existing IP CCTV infrastructure (RTSP/ONVIF).",
                "Software-based solution \u2013 no need for new hardware.",
                "Proven AI technologies: YOLO/SSD, OCR, Face Recognition, OpenCV.",
                "Public datasets (COCO, WIDER FACE, PETS, CUHK Square) for training & evaluation.",
                "Edge-ready \u2013 works on laptops, GPUs or edge devices.",
                "Modular architecture for incremental feature addition."
            ],
            "badge": "\U0001F4CA TECHNICALLY FEASIBLE, SCALABLE & COST-EFFECTIVE"
        },
        {
            "title": "POTENTIAL CHALLENGES & RISKS", "color": C_RED_DARK, "bg": C_RED_BG, "border": C_RED_BORDER, "icon": "\u26A0\uFE0F",
            "items": [
                "Poor quality / night-time video (low visibility)",
                "Multiple CCTV streams processing load",
                "False positives / false negatives in edge cases",
                "Existing legacy camera integration variance",
                "Face & number plate recognition at distance/angles",
                "Large feature scope management in Phase 1"
            ],
            "badge": None
        },
        {
            "title": "STRATEGIES TO OVERCOME", "color": C_BLUE, "bg": C_BLUE_BG, "border": C_BLUE_BORDER, "icon": "\U0001F6E1\uFE0F",
            "items": [
                "Data: Denoising, enhancement, CLAHE & day/night augmentation.",
                "Performance: Model pruning, FP16 quantization, TensorRT/GPU edge acceleration.",
                "Accuracy: Threshold tuning + DeepSORT tracking + confidence verification.",
                "Integration: Standard RTSP/ONVIF connectors with modular microservices.",
                "Privacy & Security: Role-based access, encrypted video storage.",
                "Scope: Prioritize core MVP \u2192 Detection \u2192 Tracking \u2192 Intrusion \u2192 Alerts."
            ],
            "badge": None
        }
    ]

    for cidx, col in enumerate(cols_s4):
        cx = col_start_x + cidx * (col_w + col_gap)
        
        # Column container shape
        cshape = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, col_top_y, col_w, col_h)
        cshape.fill.solid()
        cshape.fill.fore_color.rgb = col['bg']
        cshape.line.color.rgb = col['border']
        cshape.line.width = Pt(1.2)

        # Header Pill
        chb = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.05), col_top_y + Inches(0.05), col_w - Inches(0.1), Inches(0.44))
        chb.fill.solid()
        chb.fill.fore_color.rgb = col['color']
        chb.line.color.rgb = col['color']
        chtf = chb.text_frame
        chtf.margin_left = chtf.margin_right = chtf.margin_top = chtf.margin_bottom = 0
        cp = chtf.paragraphs[0]
        cp.text = f"{col['icon']}  {col['title']}"
        cp.font.name = FONT_TITLE
        cp.font.size = Pt(11)
        cp.font.bold = True
        cp.font.color.rgb = C_BG_WHITE
        cp.alignment = PP_ALIGN.CENTER

        # Content Box
        cbox = slide4.shapes.add_textbox(cx + Inches(0.12), col_top_y + Inches(0.56), col_w - Inches(0.24), col_h - Inches(1.3 if col['badge'] else 0.65))
        ctf = cbox.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0

        for itm_idx, itm in enumerate(col['items']):
            p = ctf.paragraphs[0] if itm_idx == 0 else ctf.add_paragraph()
            bullet_sym = "\u2714 " if cidx == 0 else ("\u26A0 " if cidx == 1 else "\U0001F527 ")
            p.text = f"{bullet_sym}{itm}"
            p.font.name = FONT_BODY
            p.font.size = Pt(9.5)
            p.font.color.rgb = C_DARK_TEXT
            p.space_after = Pt(7)

        # Bottom Badge for column 1
        if col['badge']:
            badge = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.15), col_top_y + col_h - Inches(0.68), col_w - Inches(0.3), Inches(0.52))
            badge.fill.solid()
            badge.fill.fore_color.rgb = C_GREEN_DARK
            badge.line.color.rgb = C_GREEN_DARK
            btf = badge.text_frame
            btf.margin_left = btf.margin_right = btf.margin_top = btf.margin_bottom = 0
            bp = btf.paragraphs[0]
            bp.text = col['badge']
            bp.font.name = FONT_TITLE
            bp.font.size = Pt(8.5)
            bp.font.bold = True
            bp.font.color.rgb = C_BG_WHITE
            bp.alignment = PP_ALIGN.CENTER

    # Bottom Banner (Navy)
    bot_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(6.05), Inches(12.533), Inches(0.85))
    bot_box.fill.solid()
    bot_box.fill.fore_color.rgb = hex_to_rgb('0F172A')
    bot_box.line.color.rgb = hex_to_rgb('1E293B')
    
    bbtf = bot_box.text_frame
    bbtf.margin_left = Inches(0.2)
    bbtf.margin_right = Inches(0.2)
    bbtf.margin_top = Inches(0.12)
    
    bbp = bbtf.paragraphs[0]
    bbp.text = "\U0001F4A1 BOTTOM LINE: Technically feasible, scalable and cost-effective \u2014"
    bbp.font.name = FONT_TITLE
    bbp.font.size = Pt(12)
    bbp.font.bold = True
    bbp.font.color.rgb = hex_to_rgb('FDE047')
    bbp.alignment = PP_ALIGN.CENTER

    bbp2 = bbtf.add_paragraph()
    bbp2.text = "with risks manageable through modular microservice design, edge optimization and phased MVP deployment."
    bbp2.font.name = FONT_BODY
    bbp2.font.size = Pt(11)
    bbp2.font.bold = True
    bbp2.font.color.rgb = C_BG_WHITE
    bbp2.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 5: Impact and Benefits
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "IMPACT AND BENEFITS", page_num=5)

    line5 = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(0.82), Inches(4.33), Inches(0.04))
    line5.fill.solid()
    line5.fill.fore_color.rgb = C_NAVY_TITLE
    line5.line.color.rgb = C_NAVY_TITLE

    # Sub-heading 1: Impact on Target Audience
    aud_title_box = slide5.shapes.add_textbox(Inches(0.4), Inches(0.95), Inches(12.533), Inches(0.3))
    attf = aud_title_box.text_frame
    attf.margin_left = attf.margin_top = attf.margin_right = attf.margin_bottom = 0
    ap = attf.paragraphs[0]
    ap.text = "\u2022  IMPACT ON TARGET AUDIENCE  \u2022"
    ap.font.name = FONT_TITLE
    ap.font.size = Pt(11.5)
    ap.font.bold = True
    ap.font.color.rgb = C_NAVY_TITLE
    ap.alignment = PP_ALIGN.CENTER

    # 5 Audience Cards
    aud_w = Inches(2.38)
    aud_gap = Inches(0.12)
    aud_start_x = Inches(0.4)
    aud_top_y = Inches(1.28)
    aud_h = Inches(2.05)

    audiences = [
        {"title": "BORDER SECURITY FORCES", "label": "DEFENSE", "color": hex_to_rgb('0284C7'), "desc": "Faster detection of suspicious activities and intrusions.\nImproved situational awareness."},
        {"title": "SECURITY OPERATORS", "label": "OPERATORS", "color": hex_to_rgb('0D9488'), "desc": "Automated detection of humans, vehicles and events.\nReduced workload and prioritized alerts."},
        {"title": "COMMAND & CONTROL CENTRES", "label": "C2 CENTRES", "color": hex_to_rgb('6366F1'), "desc": "Centralized alerts, event logs, video snippets and metadata.\nFaster decision-making and response."},
        {"title": "GOVERNMENT & AGENCIES", "label": "AGENCIES", "color": hex_to_rgb('D97706'), "desc": "Scalable solution using existing infrastructure.\nNo need for dedicated FRS/ANPR smart hardware."},
        {"title": "REMOTE BORDER POSTS", "label": "REMOTE POSTS", "color": hex_to_rgb('059669'), "desc": "AI-assisted monitoring even in resource-limited and remote perimeter locations."}
    ]

    for aidx, aud in enumerate(audiences):
        ax = aud_start_x + aidx * (aud_w + aud_gap)
        ashape = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, ax, aud_top_y, aud_w, aud_h)
        ashape.fill.solid()
        ashape.fill.fore_color.rgb = hex_to_rgb('F8FAFC')
        ashape.line.color.rgb = hex_to_rgb('CBD5E1')
        ashape.line.width = Pt(1.2)

        # Top Badge Header
        top_badge = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, ax + Inches(0.04), aud_top_y + Inches(0.04), aud_w - Inches(0.08), Inches(0.32))
        top_badge.fill.solid()
        top_badge.fill.fore_color.rgb = aud['color']
        top_badge.line.color.rgb = aud['color']
        tbtf = top_badge.text_frame
        tbtf.margin_left = tbtf.margin_right = tbtf.margin_top = tbtf.margin_bottom = 0
        tbp = tbtf.paragraphs[0]
        tbp.text = aud['title']
        tbp.font.name = FONT_TITLE
        tbp.font.size = Pt(7.8)
        tbp.font.bold = True
        tbp.font.color.rgb = C_BG_WHITE
        tbp.alignment = PP_ALIGN.CENTER

        # Audience Text
        abox = slide5.shapes.add_textbox(ax + Inches(0.08), aud_top_y + Inches(0.44), aud_w - Inches(0.16), aud_h - Inches(0.48))
        atf = abox.text_frame
        atf.word_wrap = True
        atf.margin_left = atf.margin_right = atf.margin_top = atf.margin_bottom = 0

        p = atf.paragraphs[0]
        p.text = aud['desc']
        p.font.name = FONT_BODY
        p.font.size = Pt(8.5)
        p.font.color.rgb = C_DARK_TEXT
        p.alignment = PP_ALIGN.CENTER

    # Sub-heading 2: Key Benefits
    ben_title_box = slide5.shapes.add_textbox(Inches(0.4), Inches(3.42), Inches(12.533), Inches(0.28))
    bttf = ben_title_box.text_frame
    bttf.margin_left = bttf.margin_top = bttf.margin_right = bttf.margin_bottom = 0
    bp = bttf.paragraphs[0]
    bp.text = "\u2022  KEY BENEFITS  \u2022"
    bp.font.name = FONT_TITLE
    bp.font.size = Pt(11.5)
    bp.font.bold = True
    bp.font.color.rgb = C_NAVY_TITLE
    bp.alignment = PP_ALIGN.CENTER

    # 4 Categorized Benefit Cards
    ben_w = Inches(3.02)
    ben_gap = Inches(0.15)
    ben_start_x = Inches(0.4)
    ben_top_y = Inches(3.72)
    ben_h = Inches(2.22)

    benefits = [
        {
            "title": "SOCIAL BENEFITS", "color": C_GREEN_DARK, "bg": C_GREEN_BG, "border": C_GREEN_BORDER,
            "bullets": ["Improved border safety", "Faster response to incidents", "Reduced operator workload", "Better situational awareness"]
        },
        {
            "title": "ECONOMIC BENEFITS", "color": C_BLUE, "bg": C_BLUE_BG, "border": C_BLUE_BORDER,
            "bullets": ["Reuses existing CCTV infrastructure", "Lower hardware cost & capex", "Open-source core technology", "Scalable multi-site deployment"]
        },
        {
            "title": "OPERATIONAL BENEFITS", "color": C_ORANGE_DARK, "bg": C_ORANGE_BG, "border": C_ORANGE_BORDER,
            "bullets": ["Real-time human & vehicle detection", "Face & license plate recognition", "Virtual-fence intrusion detection", "Automated alerts & evidence logging"]
        },
        {
            "title": "ENVIRONMENTAL / RESOURCE", "color": C_TEAL, "bg": C_TEAL_BG, "border": C_TEAL_BORDER,
            "bullets": ["Extends operational life of existing cameras", "Reduces unnecessary e-waste & replacement", "Edge processing minimizes cloud bandwidth", "Energy-efficient compute architecture"]
        }
    ]

    for bidx, ben in enumerate(benefits):
        bx = ben_start_x + bidx * (ben_w + ben_gap)
        bshape = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, ben_top_y, ben_w, ben_h)
        bshape.fill.solid()
        bshape.fill.fore_color.rgb = ben['bg']
        bshape.line.color.rgb = ben['border']
        bshape.line.width = Pt(1.2)

        # Header Pill
        bhb = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx + Inches(0.04), ben_top_y + Inches(0.04), ben_w - Inches(0.08), Inches(0.36))
        bhb.fill.solid()
        bhb.fill.fore_color.rgb = ben['color']
        bhb.line.color.rgb = ben['color']
        bhtf = bhb.text_frame
        bhtf.margin_left = bhtf.margin_right = bhtf.margin_top = bhtf.margin_bottom = 0
        p = bhtf.paragraphs[0]
        p.text = ben['title']
        p.font.name = FONT_TITLE
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = C_BG_WHITE
        p.alignment = PP_ALIGN.CENTER

        # Bullets
        bbox = slide5.shapes.add_textbox(bx + Inches(0.08), ben_top_y + Inches(0.44), ben_w - Inches(0.16), ben_h - Inches(0.48))
        btf = bbox.text_frame
        btf.word_wrap = True
        btf.margin_left = btf.margin_right = btf.margin_top = btf.margin_bottom = 0
        for itm_idx, itm in enumerate(ben['bullets']):
            p = btf.paragraphs[0] if itm_idx == 0 else btf.add_paragraph()
            p.text = f"\u2022 {itm}"
            p.font.name = FONT_BODY
            p.font.size = Pt(8.8)
            p.font.color.rgb = C_DARK_TEXT
            p.space_after = Pt(3)

    # Bottom Overall Impact Banner
    o_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(6.05), Inches(12.533), Inches(0.85))
    o_box.fill.solid()
    o_box.fill.fore_color.rgb = hex_to_rgb('064E3B')
    o_box.line.color.rgb = hex_to_rgb('047857')

    otf = o_box.text_frame
    otf.margin_left = Inches(0.2)
    otf.margin_right = Inches(0.2)
    otf.margin_top = Inches(0.12)
    op = otf.paragraphs[0]
    op.text = "\U0001F3AF OVERALL IMPACT: EXISTING CCTV + AI INTELLIGENCE + REAL-TIME ALERTS"
    op.font.name = FONT_TITLE
    op.font.size = Pt(12)
    op.font.bold = True
    op.font.color.rgb = hex_to_rgb('FDE047')
    op.alignment = PP_ALIGN.CENTER

    op2 = otf.add_paragraph()
    op2.text = "= A SMARTER, FASTER & MORE PROACTIVE BORDER-SECURITY NETWORK"
    op2.font.name = FONT_TITLE
    op2.font.size = Pt(11.5)
    op2.font.bold = True
    op2.font.color.rgb = C_BG_WHITE
    op2.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 6: Research and References
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "RESEARCH AND REFERENCES", page_num=6)

    line6 = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.2), Inches(0.82), Inches(4.93), Inches(0.04))
    line6.fill.solid()
    line6.fill.fore_color.rgb = C_NAVY_TITLE
    line6.line.color.rgb = C_NAVY_TITLE

    # Row 1: 5 Pillar Cards
    rc_w = Inches(2.38)
    rc_gap = Inches(0.12)
    rc_start_x = Inches(0.4)
    rc_top_y = Inches(1.05)
    rc_h = Inches(3.8)

    r_cards = [
        {
            "num": "1. PROBLEM STATEMENT", "color": C_BLUE, "bg": C_BLUE_BG, "border": C_BLUE_BORDER,
            "type": "ps",
            "ps_id": "SIH26187",
            "ps_title": "AI-Based Intelligent Video Analytics Platform for Border Surveillance using Existing CCTV Infrastructure",
            "org": "Ministry of Home Affairs",
            "cat": "Smart Automation",
            "link": "sih2026.vuce.in/en/ps/SIH26187"
        },
        {
            "num": "2. TECHNOLOGIES USED", "color": C_GREEN_DARK, "bg": C_GREEN_BG, "border": C_GREEN_BORDER,
            "type": "tech",
            "items": [
                ("OpenCV", "Image & Video Processing"),
                ("YOLO / SSD", "Object Detection"),
                ("PyTorch", "Deep Learning Framework"),
                ("OCR", "License Plate Recognition"),
                ("Face Recognition", "Identity Matching")
            ]
        },
        {
            "num": "3. DATASETS", "color": C_ORANGE_DARK, "bg": C_ORANGE_BG, "border": C_ORANGE_BORDER,
            "type": "tech",
            "items": [
                ("COCO", "Object Detection"),
                ("WIDER FACE", "Face Detection"),
                ("PETS / CUHK Square", "Surveillance Videos"),
                ("OpenALPR", "License Plate Recognition")
            ]
        },
        {
            "num": "4. EVALUATION METRICS", "color": C_PURPLE, "bg": C_PURPLE_BG, "border": C_PURPLE_BORDER,
            "type": "metrics",
            "items": [
                "mAP (Mean Average Precision)",
                "Accuracy & Precision",
                "Recall & F1 Score",
                "AUC-ROC Curve",
                "False Alarm Rate (FAR)",
                "Processing Latency (ms)",
                "Frames Per Second (FPS)"
            ]
        },
        {
            "num": "5. DEPLOYMENT & TOOLS", "color": C_TEAL, "bg": C_TEAL_BG, "border": C_TEAL_BORDER,
            "type": "tech",
            "items": [
                ("FastAPI", "Backend APIs"),
                ("PostgreSQL / SQLite", "Data Storage"),
                ("NVIDIA TensorRT", "Inference Optimization"),
                ("Edge / GPU Deployment", "Real-Time Processing")
            ]
        }
    ]

    for idx, rc in enumerate(r_cards):
        rx = rc_start_x + idx * (rc_w + rc_gap)
        rshape = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx, rc_top_y, rc_w, rc_h)
        rshape.fill.solid()
        rshape.fill.fore_color.rgb = rc['bg']
        rshape.line.color.rgb = rc['border']
        rshape.line.width = Pt(1.2)

        # Header Pill
        rhb = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx + Inches(0.04), rc_top_y + Inches(0.04), rc_w - Inches(0.08), Inches(0.38))
        rhb.fill.solid()
        rhb.fill.fore_color.rgb = rc['color']
        rhb.line.color.rgb = rc['color']
        rhtf = rhb.text_frame
        rhtf.margin_left = rhtf.margin_right = rhtf.margin_top = rhtf.margin_bottom = 0
        p = rhtf.paragraphs[0]
        p.text = rc['num']
        p.font.name = FONT_TITLE
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = C_BG_WHITE
        p.alignment = PP_ALIGN.CENTER

        # Content inside
        rbox = slide6.shapes.add_textbox(rx + Inches(0.06), rc_top_y + Inches(0.44), rc_w - Inches(0.12), rc_h - Inches(0.48))
        rtf = rbox.text_frame
        rtf.word_wrap = True
        rtf.margin_left = rtf.margin_right = rtf.margin_top = rtf.margin_bottom = 0

        if rc['type'] == 'ps':
            p = rtf.paragraphs[0]
            p.text = "Smart India Hackathon 2026"
            p.font.name = FONT_BODY
            p.font.size = Pt(8.5)
            p.font.color.rgb = C_DARK_TEXT
            p.alignment = PP_ALIGN.CENTER

            p_id = rtf.add_paragraph()
            p_id.text = rc['ps_id']
            p_id.font.name = FONT_TITLE
            p_id.font.size = Pt(10)
            p_id.font.bold = True
            p_id.font.color.rgb = C_BLUE
            p_id.alignment = PP_ALIGN.CENTER
            p_id.space_after = Pt(2)

            p_title = rtf.add_paragraph()
            p_title.text = rc['ps_title']
            p_title.font.name = FONT_BODY
            p_title.font.size = Pt(7.8)
            p_title.font.bold = True
            p_title.font.color.rgb = C_DARK_TEXT
            p_title.alignment = PP_ALIGN.CENTER
            p_title.space_after = Pt(4)

            p_org = rtf.add_paragraph()
            p_org.text = f"{rc['org']}\n{rc['cat']}"
            p_org.font.name = FONT_BODY
            p_org.font.size = Pt(7.8)
            p_org.font.color.rgb = C_MUTED_TEXT
            p_org.alignment = PP_ALIGN.CENTER
            p_org.space_after = Pt(4)

            p_lnk = rtf.add_paragraph()
            p_lnk.text = rc['link']
            p_lnk.font.name = FONT_BODY
            p_lnk.font.size = Pt(7.2)
            p_lnk.font.color.rgb = C_BLUE
            p_lnk.alignment = PP_ALIGN.CENTER

        elif rc['type'] == 'tech':
            for tidx, (tname, tdesc) in enumerate(rc['items']):
                p = rtf.paragraphs[0] if tidx == 0 else rtf.add_paragraph()
                p.text = f"\u2022 {tname}"
                p.font.name = FONT_TITLE
                p.font.size = Pt(8.5)
                p.font.bold = True
                p.font.color.rgb = C_DARK_TEXT
                
                p2 = rtf.add_paragraph()
                p2.text = f"   ({tdesc})"
                p2.font.name = FONT_BODY
                p2.font.size = Pt(7.5)
                p2.font.color.rgb = C_MUTED_TEXT
                p2.space_after = Pt(3)

        elif rc['type'] == 'metrics':
            for midx, mitem in enumerate(rc['items']):
                p = rtf.paragraphs[0] if midx == 0 else rtf.add_paragraph()
                p.text = f"\u2022 {mitem}"
                p.font.name = FONT_BODY
                p.font.size = Pt(8.2)
                p.font.bold = True if midx < 3 else False
                p.font.color.rgb = C_DARK_TEXT
                p.space_after = Pt(3)

    # Row 2 (Bottom Pipeline): Research -> Development Pipeline
    pipe_y = Inches(5.0)
    pipe_h = Inches(1.95)
    
    pipe_shape = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), pipe_y, Inches(12.533), pipe_h)
    pipe_shape.fill.solid()
    pipe_shape.fill.fore_color.rgb = hex_to_rgb('F8FAFC')
    pipe_shape.line.color.rgb = hex_to_rgb('CBD5E1')
    pipe_shape.line.width = Pt(1.2)

    # Pipeline Title
    pt_box = slide6.shapes.add_textbox(Inches(0.55), pipe_y + Inches(0.06), Inches(12.2), Inches(0.32))
    pttf = pt_box.text_frame
    pttf.margin_left = pttf.margin_top = pttf.margin_right = pttf.margin_bottom = 0
    p = pttf.paragraphs[0]
    p.text = "RESEARCH  \u2192  DEVELOPMENT PIPELINE"
    p.font.name = FONT_TITLE
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_TITLE

    # 6 Step Pipeline Boxes
    steps_r6 = [
        ("Step 1", "Research Papers\n& Datasets", hex_to_rgb('0284C7')),
        ("Step 2", "AI Models\n(Train & Test)", hex_to_rgb('7C3AED')),
        ("Step 3", "Computer Vision\n& Analytics", hex_to_rgb('10B981')),
        ("Step 4", "ANPR / Face\nRecognition", hex_to_rgb('D97706')),
        ("Step 5", "Testing & Eval\n(Metrics)", hex_to_rgb('EC4899')),
        ("Step 6", "Real-Time\nDeployment", hex_to_rgb('2563EB'))
    ]

    pstep_w = Inches(1.8)
    pstep_gap = Inches(0.24)
    pstep_start_x = Inches(0.55)
    pstep_y = pipe_y + Inches(0.42)
    pstep_h = Inches(1.38)

    for psidx, (psnum, pstitle, pscolor) in enumerate(steps_r6):
        psx = pstep_start_x + psidx * (pstep_w + pstep_gap)
        psbox = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, psx, pstep_y, pstep_w, pstep_h)
        psbox.fill.solid()
        psbox.fill.fore_color.rgb = C_BG_WHITE
        psbox.line.color.rgb = pscolor
        psbox.line.width = Pt(1.5)

        # Step Header Pill
        s_pill = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, psx + Inches(0.04), pstep_y + Inches(0.04), pstep_w - Inches(0.08), Inches(0.32))
        s_pill.fill.solid()
        s_pill.fill.fore_color.rgb = pscolor
        s_pill.line.color.rgb = pscolor
        sptf = s_pill.text_frame
        sptf.margin_left = sptf.margin_right = sptf.margin_top = sptf.margin_bottom = 0
        p = sptf.paragraphs[0]
        p.text = psnum
        p.font.name = FONT_TITLE
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = C_BG_WHITE
        p.alignment = PP_ALIGN.CENTER

        # Step Text
        stext = slide6.shapes.add_textbox(psx + Inches(0.04), pstep_y + Inches(0.44), pstep_w - Inches(0.08), pstep_h - Inches(0.48))
        stf = stext.text_frame
        stf.word_wrap = True
        stf.margin_left = stf.margin_right = stf.margin_top = stf.margin_bottom = 0
        
        p = stf.paragraphs[0]
        p.text = pstitle
        p.font.name = FONT_TITLE
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_TITLE
        p.alignment = PP_ALIGN.CENTER

        # Arrow to next step
        if psidx < len(steps_r6) - 1:
            arr_box = slide6.shapes.add_textbox(psx + pstep_w, pstep_y + Inches(0.45), pstep_gap, Inches(0.4))
            atf = arr_box.text_frame
            atf.margin_left = atf.margin_right = atf.margin_top = atf.margin_bottom = 0
            p = atf.paragraphs[0]
            p.text = "\u2192"
            p.font.name = FONT_TITLE
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = hex_to_rgb('64748B')
            p.alignment = PP_ALIGN.CENTER

    # Save presentation
    output_path = r'c:\Users\FALCON JNB\.gemini\antigravity-ide\scratch\drishtix\DRISHTIX_SIH2026_Fully_Editable_Presentation.pptx'
    prs.save(output_path)
    print(f'Successfully generated PowerPoint presentation at: {output_path}')

    # Also save a copy to Downloads folder for convenient access
    downloads_path = os.path.expanduser('~\\Downloads\\DRISHTIX_SIH2026_Fully_Editable_Presentation.pptx')
    try:
        prs.save(downloads_path)
        print(f'Also saved copy to Downloads: {downloads_path}')
    except Exception as e:
        print('Could not save to Downloads:', e)

if __name__ == '__main__':
    create_presentation()
